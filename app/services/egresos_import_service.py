"""Lectura acotada de gastos XLSX; la vista previa no escribe movimientos."""

import base64
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import hashlib
import hmac
from io import BytesIO
import json
import re
import time
from typing import Literal
import unicodedata
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel
from pydantic import BaseModel, Field, model_validator

from app.core.config import settings

MAX_BYTES = 5 * 1024 * 1024
MAX_EXPANDED_BYTES = 25 * 1024 * 1024
MAX_ROWS = 500
MAX_SCAN_ROWS = 5000
MAX_COLUMNS = 40
TOKEN_TTL = 3600
CAMPOS = ("fecha", "proveedor", "monto", "metodo_pago", "concepto", "categoria", "notas")
OBLIGATORIOS = CAMPOS[:4]
ALIAS = {
    "fecha": {"fecha", "dia", "fechagasto", "fechadepago", "fechadecompra"},
    "proveedor": {"proveedor", "proveedorpersona", "proveedoropersona", "establecimiento", "comercio", "beneficiario"},
    "monto": {"monto", "importe", "total", "gasto", "importecompras", "cantidad", "valor"},
    "metodo_pago": {"pago", "formadepago", "formapago", "metododepago", "metodopago", "tipodepago"},
    "concepto": {"concepto", "descripcion", "detalle", "motivo"},
    "categoria": {"categoria", "rubro", "tipodegasto"},
    "notas": {"notas", "observaciones", "referencia"},
}
MESES = {m: i + 1 for i, m in enumerate(("ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"))}
MESES.update({"jan": 1, "apr": 4, "aug": 8, "dec": 12})


class OpcionesImportacion(BaseModel):
    hoja: str | None = Field(default=None, max_length=100)
    fila_encabezado: int | None = Field(default=None, ge=0, le=50)
    columnas: dict[str, int | None] | None = None
    categoria: str = Field(default="operativo", max_length=60)
    formato_fecha: Literal["dia_mes_anio", "mes_dia_anio"] = "dia_mes_anio"
    guardar_proveedor: bool = True

    @model_validator(mode="after")
    def validar_columnas(self):
        if self.columnas is not None:
            if set(self.columnas) - set(CAMPOS):
                raise ValueError("Columna de destino desconocida")
            for valor in self.columnas.values():
                if valor is not None and (isinstance(valor, bool) or not 1 <= valor <= MAX_COLUMNS):
                    raise ValueError("Indice de columna invalido")
        return self


def normalizar(valor) -> str:
    texto = unicodedata.normalize("NFKD", str(valor or "")).casefold()
    return "".join(c for c in texto if c.isalnum() and not unicodedata.combining(c))


def texto(valor) -> str:
    return re.sub(r"\s+", " ", str(valor if valor is not None else "")).strip()


def fecha_excel(valor, epoch, formato="dia_mes_anio") -> str:
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        convertido = from_excel(valor, epoch)
        if isinstance(convertido, datetime):
            return convertido.date().isoformat()
        raise ValueError("Fecha de Excel invalida")
    raw = texto(valor).lower().rstrip(".")
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", raw):
        try:
            return date.fromisoformat(raw).isoformat()
        except ValueError:
            raise ValueError("Fecha invalida; revisa el dia, mes y ano") from None
    partes = re.fullmatch(r"(\d{1,2})[\s/.-]+([a-záéíóú]+|\d{1,2})[\s/.-]+(\d{2}|\d{4})", raw)
    if not partes:
        raise ValueError("Fecha invalida; usa dia/mes/ano")
    dia, mes, anio = partes.groups()
    if mes.isdigit() and formato == "mes_dia_anio":
        dia, mes = mes, dia
    mes = int(mes) if mes.isdigit() else MESES.get(normalizar(mes)[:3], 0)
    anio = int(anio) + (2000 if len(anio) == 2 else 0)
    try:
        return date(anio, mes, int(dia)).isoformat()
    except ValueError:
        raise ValueError("Fecha invalida; revisa el dia, mes y ano") from None


def monto_excel(valor) -> str:
    if isinstance(valor, bool) or valor is None:
        raise ValueError("Monto obligatorio")
    raw = texto(valor)
    raw = re.sub(r"(?i)MXN|\$|\s", "", raw)
    if "," in raw and "." in raw:
        if raw.rfind(",") > raw.rfind("."):
            if not re.fullmatch(r"[+-]?\d{1,3}(\.\d{3})+,\d+", raw):
                raise ValueError("Separadores del monto invalidos")
            raw = raw.replace(".", "").replace(",", ".")
        else:
            if not re.fullmatch(r"[+-]?\d{1,3}(,\d{3})+\.\d+", raw):
                raise ValueError("Separadores del monto invalidos")
            raw = raw.replace(",", "")
    elif "," in raw:
        if re.fullmatch(r"[+-]?\d{1,3}(,\d{3})+", raw):
            raw = raw.replace(",", "")
        elif re.fullmatch(r"[+-]?\d+,\d{1,2}", raw):
            raw = raw.replace(",", ".")
        else:
            raise ValueError("Separadores del monto invalidos")
    try:
        if isinstance(valor, str) and "." in raw and len(raw.rsplit(".", 1)[1]) > 2:
            raise ValueError("Monto ambiguo o con mas de dos decimales")
        monto = Decimal(raw)
        if not monto.is_finite() or not 0 < monto <= Decimal("10000000"):
            raise ValueError("Monto fuera de rango")
        centavos = monto.quantize(Decimal("0.01"))
        if abs(monto - centavos) > Decimal("0.00000001"):
            raise ValueError("El monto debe tener hasta dos decimales")
        return format(centavos, ".2f")
    except InvalidOperation:
        raise ValueError("Monto invalido; solo se admiten pesos MXN") from None


def metodo_excel(valor) -> str:
    clave = normalizar(valor)
    return {
        "efectivo": "efectivo", "cash": "efectivo", "transferencia": "transferencia",
        "transferenciabancaria": "transferencia", "spei": "transferencia", "bbva": "bbva",
        "clip": "clip", "tarjeta": "tarjeta", "credito": "credito", "acredito": "credito",
        "debito": "debito", "mixto": "mixto", "resguardo": "resguardo",
    }.get(clave, texto(valor).lower())


def _columnas(fila) -> dict:
    return {campo: next((i + 1 for i, celda in enumerate(fila)
                        if normalizar(celda.value) in alias), None) for campo, alias in ALIAS.items()}


def leer_excel(contenido: bytes, opciones: OpcionesImportacion) -> dict:
    if len(contenido) > MAX_BYTES:
        raise ValueError("El Excel supera 5 MB")
    try:
        with ZipFile(BytesIO(contenido)) as archivo:
            miembros = archivo.infolist()
            if (len(miembros) > 1000 or sum(m.file_size for m in miembros) > MAX_EXPANDED_BYTES
                    or any(m.flag_bits & 1 or "vbaproject" in m.filename.lower() for m in miembros)):
                raise ValueError("Excel demasiado grande, cifrado o con macros")
            if "xl/workbook.xml" not in archivo.namelist():
                raise ValueError("El archivo no es un libro XLSX")
        wb = load_workbook(BytesIO(contenido), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, KeyError, OSError):
        raise ValueError("No se pudo abrir el Excel; usa un archivo .xlsx sin contrasena") from None
    try:
        hojas = wb.sheetnames
        if not hojas or len(hojas) > 50:
            raise ValueError("El Excel debe tener entre 1 y 50 hojas")
        hoja = opciones.hoja or hojas[0]
        if hoja not in hojas:
            raise ValueError("La hoja seleccionada no existe")
        ws = wb[hoja]
        ws.reset_dimensions()
        filas = []
        for numero, fila in enumerate(ws.iter_rows(max_col=MAX_COLUMNS), 1):
            if numero > MAX_SCAN_ROWS:
                raise ValueError("La hoja supera 5000 filas; importa por partes")
            filas.append(fila)
        if not any(any(c.value is not None for c in fila) for fila in filas):
            raise ValueError("La hoja seleccionada esta vacia")
        encabezado = opciones.fila_encabezado
        if encabezado is None:
            mejor = max(range(min(30, len(filas))), key=lambda i: sum(v is not None for v in _columnas(filas[i]).values()))
            encabezado = mejor + 1 if sum(v is not None for v in _columnas(filas[mejor]).values()) >= 2 else 0
        if encabezado > len(filas):
            raise ValueError("La fila de encabezados no existe")
        titulos = filas[encabezado - 1] if encabezado else []
        columnas = opciones.columnas if opciones.columnas is not None else (
            _columnas(titulos) if titulos else dict(zip(OBLIGATORIOS, range(1, 5), strict=True))
        )
        columnas = {campo: columnas.get(campo) for campo in CAMPOS}
        faltantes = [campo for campo in OBLIGATORIOS if columnas[campo] is None]
        elegidas = [c for c in columnas.values() if c is not None]
        if len(set(elegidas)) != len(elegidas):
            raise ValueError("No asignes la misma columna a dos campos")
        ancho = max((i + 1 for fila in filas for i, c in enumerate(fila) if c.value is not None), default=4)
        salida, registros = [], 0
        for numero, fila in enumerate(filas[encabezado:], encabezado + 1):
            valores = [c.value for c in fila]
            if not any(v is not None and texto(v) for v in valores):
                continue
            raw = {campo: fila[idx - 1].value if idx else None for campo, idx in columnas.items()}
            errores, datos = [], {}
            primero = next((normalizar(v) for v in valores if texto(v)), "")
            es_total = primero in {"total", "totales", "subtotal", "totalgastos", "totaldegastos", "sumatotal"}
            if not es_total:
                registros += 1
                if registros > MAX_ROWS:
                    raise ValueError("Maximo 500 gastos por importacion; divide el archivo")
            if not es_total and not faltantes:
                try:
                    if any(fila[idx - 1].data_type in {"f", "e"} for idx in elegidas):
                        raise ValueError("Hay formulas o errores de Excel; convierte esa fila a valores")
                    if not texto(raw["metodo_pago"]):
                        raise ValueError("Forma de pago obligatoria; no se puede suponer efectivo")
                    datos = {
                        "fecha": fecha_excel(raw["fecha"], wb.epoch, opciones.formato_fecha),
                        "proveedor": texto(raw["proveedor"]),
                        "monto": monto_excel(raw["monto"]),
                        "metodo_pago": metodo_excel(raw["metodo_pago"]),
                        "concepto": texto(raw["concepto"]) or texto(raw["proveedor"]),
                        "categoria": normalizar(raw["categoria"] or opciones.categoria),
                        "notas": texto(raw["notas"]) or None,
                    }
                except (ValueError, TypeError, OverflowError) as exc:
                    errores.append(str(exc) or "Dato invalido")
            salida.append({"fila": numero, "datos": datos, "errores": errores,
                           "omitida": es_total, "original": {k: texto(v)[:200] for k, v in raw.items()}})
        return {"hojas": hojas, "hoja": hoja, "fila_encabezado": encabezado, "columnas": columnas,
                "opciones_columnas": [{"indice": i + 1, "nombre": get_column_letter(i + 1) + (
                    " - " + texto(titulos[i].value)[:80] if titulos and titulos[i].value is not None else "")}
                    for i in range(ancho)], "faltantes": faltantes, "filas": salida}
    finally:
        wb.close()


def clave_fila(huella: str, hoja: str, fila: int) -> str:
    return "excel-" + hashlib.sha256(json.dumps([huella, hoja, fila], ensure_ascii=True).encode()).hexdigest()


def firmar_vista_previa(datos: dict, usuario_id: int) -> str:
    payload = {**datos, "usuario_id": usuario_id, "expira": int(time.time()) + TOKEN_TTL}
    raw = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode()
    mensaje = base64.urlsafe_b64encode(raw).decode()
    firma = hmac.new(settings.SECRET_KEY.encode(), b"egresos-excel-v1:" + mensaje.encode(), hashlib.sha256).hexdigest()
    return mensaje + "." + firma


def leer_vista_previa(token: str, usuario_id: int) -> dict:
    try:
        mensaje, firma = token.split(".")
        esperada = hmac.new(settings.SECRET_KEY.encode(), b"egresos-excel-v1:" + mensaje.encode(), hashlib.sha256).hexdigest()
        if not hmac.compare_digest(firma, esperada):
            raise ValueError
        datos = json.loads(base64.urlsafe_b64decode(mensaje))
        if datos["usuario_id"] != usuario_id or datos["expira"] < time.time():
            raise ValueError
        return datos
    except (ValueError, KeyError, TypeError):
        raise ValueError("La vista previa vencio o cambio; vuelve a leer el Excel") from None

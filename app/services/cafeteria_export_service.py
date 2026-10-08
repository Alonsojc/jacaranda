"""Customer statements and complete, read-only cafeteria ledger exports."""

from datetime import datetime
from decimal import Decimal
from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.orm import Session

from app.core.time_utils import operation_datetime, operation_now
from app.models.cafeteria import CafeteriaCliente, EstadoCuentaCafeteria
from app.services import cafeteria_service as svc
from app.services.excel_service import _auto_width, _style_header_row
from app.services.pdf_service import _header


def _metodo(pago) -> str:
    if pago.terminal.value in {"clip", "bbva"}:
        return pago.terminal.value.upper()
    return {"01": "Efectivo", "03": "Transferencia", "04": "Tarjeta crédito",
            "28": "Tarjeta débito"}.get(pago.metodo_pago.value, pago.metodo_pago.value)


def _cantidad(value) -> str:
    return format(Decimal(value).normalize(), "f")


def _saldo(venta) -> Decimal:
    return Decimal("0") if venta.estado == EstadoCuentaCafeteria.CANCELADA else venta.saldo_pendiente


def estado_cuenta(db: Session, cafeteria_id=None, cafeteria_nombre=None) -> dict:
    if cafeteria_id is None and not cafeteria_nombre:
        raise ValueError("Selecciona una cafetería para su estado de cuenta")
    if cafeteria_id is not None and cafeteria_nombre is not None:
        raise ValueError("Selecciona una sola cafetería")
    if cafeteria_id is not None:
        cliente = db.get(CafeteriaCliente, cafeteria_id)
        if not cliente:
            raise ValueError("Cafetería no encontrada")
        nombre = cliente.nombre
    else:
        anteriores = svc.listar_ventas(db, limit=1, cafeteria_nombre=cafeteria_nombre)
        if not anteriores:
            raise ValueError("Cafetería no encontrada")
        nombre = anteriores[0].cafeteria_nombre
    ventas = svc.listar_ventas(db, limit=None, cafeteria_id=cafeteria_id,
                              cafeteria_nombre=cafeteria_nombre, pendientes=True)
    return {
        "cliente": nombre, "generado_en": operation_now(), "ventas": ventas,
        "total": sum((v.total for v in ventas), Decimal("0")),
        "pagado": sum((v.monto_pagado for v in ventas), Decimal("0")),
        "saldo": sum((_saldo(v) for v in ventas), Decimal("0")),
    }


def generar_estado_cuenta_pdf(db: Session, cafeteria_id=None, cafeteria_nombre=None) -> BytesIO:
    data = estado_cuenta(db, cafeteria_id, cafeteria_nombre)
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=17 * mm, bottomMargin=18 * mm)
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle("CuentaCell", parent=styles["Normal"], fontSize=8, leading=11)
    small_style = ParagraphStyle("CuentaSmall", parent=styles["Normal"], fontSize=8, leading=11,
                                textColor=colors.HexColor("#555555"))
    story = []
    _header(story, styles, "Estado de cuenta de Cafeterías", escape(data["cliente"]))
    story.append(Paragraph(f"Al {data['generado_en']:%d/%m/%Y %H:%M} | Importes en MXN", small_style))
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph(f"<b>Saldo por pagar: ${data['saldo']:,.2f}</b>", styles["Heading2"]))
    story.append(Paragraph(f"{len(data['ventas'])} pedidos con saldo | Total: ${data['total']:,.2f}"
                           f" | Abonos: ${data['pagado']:,.2f}", styles["Normal"]))
    story.append(Spacer(1, 5 * mm))
    if not data["ventas"]:
        story.append(Paragraph("No hay pedidos pendientes de pago.", styles["Normal"]))
    else:
        rows = [["Pedido / productos", "Entregado", "Vence", "Total", "Abonos", "Saldo"]]
        for venta in reversed(data["ventas"]):
            productos = "<br/>".join(
                escape(f"{_cantidad(d.cantidad)}x {d.producto_nombre or 'Producto'}")
                for d in venta.detalles
            )
            entrega = venta.fecha_entrega.strftime("%d/%m/%Y") if venta.fecha_entrega else "No registrada"
            rows.append([
                Paragraph(f"<b>{escape(venta.folio)}</b><br/>{productos}", cell_style),
                Paragraph(entrega, cell_style),
                venta.fecha_credito.strftime("%d/%m/%Y") if venta.fecha_credito else "",
                f"${venta.total:,.2f}", f"${venta.monto_pagado:,.2f}", f"${_saldo(venta):,.2f}",
            ])
        rows.append(["TOTAL", "", "", f"${data['total']:,.2f}",
                     f"${data['pagado']:,.2f}", f"${data['saldo']:,.2f}"])
        table = Table(rows, colWidths=[61 * mm, 23 * mm, 23 * mm, 24 * mm, 24 * mm, 24 * mm],
                      repeatRows=1, splitInRow=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#c4988a")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8), ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ALIGN", (3, 1), (-1, -1), "RIGHT"),
            ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#dddddd")),
            ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ]))
        story.append(table)
        pagos = [(venta.folio, pago) for venta in data["ventas"] for pago in venta.pagos]
        if pagos:
            story.append(Spacer(1, 6 * mm))
            story.append(Paragraph("Abonos de los pedidos pendientes", styles["Heading3"]))
            for folio, pago in sorted(pagos, key=lambda x: (x[1].fecha, x[1].id)):
                fecha = operation_datetime(pago.fecha).strftime("%d/%m/%Y")
                texto = f"{folio} | {fecha} | {_metodo(pago)} | ${pago.monto:,.2f}"
                if pago.referencia:
                    texto += f" | {pago.referencia}"
                story.append(Paragraph(escape(texto), small_style))

    def footer(canvas, _doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawString(18 * mm, 10 * mm, "Jacaranda - Estado de cuenta | MXN")
        canvas.drawRightString(letter[0] - 18 * mm, 10 * mm, f"Página {_doc.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    buf.seek(0)
    return buf


def exportar_historial_excel(db: Session) -> BytesIO:
    ventas = svc.listar_ventas(db, limit=None)
    wb = Workbook()
    pedidos = wb.active
    pedidos.title = "Pedidos"
    pedidos.append(["Pedido ID", "Folio", "Cliente ID", "Cliente", "Contacto", "Teléfono",
                    "Registrado", "Entregado", "Vencimiento", "Días crédito", "Estado",
                    "Subtotal", "IVA", "Total", "Pagado", "Saldo por cobrar", "Último pago",
                    "Notas", "Usuario ID"])
    partidas = wb.create_sheet("Partidas")
    partidas.append(["Partida ID", "Pedido ID", "Folio", "Cliente", "Producto ID", "Producto",
                     "Cantidad", "Precio unitario", "Subtotal", "IVA", "Total"])
    pagos = wb.create_sheet("Pagos")
    pagos.append(["Pago ID", "Pedido ID", "Folio", "Cliente", "Estado pedido", "Fecha pago",
                  "Monto", "Método", "Código método", "Terminal", "Referencia", "Usuario ID"])
    for venta in reversed(ventas):
        ultimo_pago = max((operation_datetime(p.fecha).date() for p in venta.pagos), default=None)
        pedidos.append([
            venta.id, venta.folio, venta.cafeteria_id, venta.cafeteria_nombre,
            venta.contacto_nombre, venta.telefono, operation_datetime(venta.fecha).replace(tzinfo=None),
            venta.fecha_entrega, venta.fecha_credito, venta.dias_credito, venta.estado.value,
            venta.subtotal, venta.total_impuestos, venta.total, venta.monto_pagado, _saldo(venta),
            ultimo_pago, venta.notas, venta.usuario_id,
        ])
        for detalle in venta.detalles:
            partidas.append([
                detalle.id, venta.id, venta.folio, venta.cafeteria_nombre, detalle.producto_id,
                detalle.producto_nombre, detalle.cantidad, detalle.precio_unitario,
                detalle.subtotal, detalle.monto_iva, detalle.subtotal + detalle.monto_iva,
            ])
        for pago in sorted(venta.pagos, key=lambda p: (p.fecha, p.id)):
            pagos.append([
                pago.id, venta.id, venta.folio, venta.cafeteria_nombre, venta.estado.value,
                operation_datetime(pago.fecha).date(), pago.monto, _metodo(pago),
                pago.metodo_pago.value, pago.terminal.value, pago.referencia, pago.usuario_id,
            ])
    for ws, money_columns in ((pedidos, range(12, 17)), (partidas, range(8, 12)), (pagos, [7])):
        _style_header_row(ws, 1, ws.max_column)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 32
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                # Customer-entered strings are data, never Excel formulas.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                if cell.row > 1 and cell.column in money_columns:
                    cell.number_format = '"$"#,##0.00'
                if isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm"
        _auto_width(ws)
        for col in ("G", "H", "I", "Q") if ws == pedidos else (["F"] if ws == pagos else []):
            for cell in ws[col][1:]:
                if cell.value is not None and not isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd"
    partidas.column_dimensions["G"].width = 14
    for cell in partidas["G"][1:]:
        cell.number_format = "0.####"
    wb.properties.title = "Historial completo de Cafeterías"
    wb.properties.subject = "Todos los pedidos, partidas y pagos. Importes en MXN."
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

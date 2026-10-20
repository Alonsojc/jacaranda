"""Importaciones con libros sinteticos, sin tocar movimientos operativos."""

from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
import json
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook
import pytest

from app.core.time_utils import operation_today
from app.models.auditoria import LogAuditoria
from app.models.egreso import Egreso
from app.models.inventario import Proveedor
from app.models.resguardo import ResguardoControl
from app.models.usuario import RolUsuario
from app.services import egresos_import_service as svc

HEADERS = ["FECHA", "PROVEEDOR", "IMPORTE", "FORMA DE PAGO"]
PATH = "/api/v1/egresos/importacion/"


def libro(filas, *, headers=HEADERS, epoch=None, segunda=None):
    wb = Workbook()
    wb.active.title = "Septiembre"
    if epoch:
        wb.epoch = epoch
    if headers:
        wb.active.append(headers)
    for fila in filas:
        wb.active.append(fila)
    if segunda:
        hoja = wb.create_sheet("Octubre")
        hoja.append(HEADERS)
        for fila in segunda:
            hoja.append(fila)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def preview(client, headers, contenido, opciones=None, nombre="gastos.xlsx"):
    response = client.post(PATH + "previsualizar", headers=headers,
        files={"archivo": (nombre, contenido, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"opciones": json.dumps(opciones or {})})
    assert response.status_code == 200, response.text
    return response.json()


def confirmar(client, headers, vista, filas=None, **extra):
    return client.post(PATH + "confirmar", headers=headers, json={
        "token": vista["token"], "filas": filas or [f["fila"] for f in vista["filas"] if f["estado"] == "lista"], **extra})


def test_preview_no_escribe_y_confirmacion_audita_reutiliza_y_no_duplica(client, auth_headers, db):
    contenido = libro([[date(2026, 9, 1), "Proveedor prueba", Decimal("178.49"), "TARJETA"],
                      ["02-sep-26", "Proveedor prueba", "$1,000.00", "EFECTIVO"],
                      ["03/09/2026", "Molino prueba", "1.300,50", "A CREDITO"],
                      ["TOTAL", None, "=SUM(C2:C4)", None]])
    vista = preview(client, auth_headers, contenido)
    assert db.query(Egreso).count() == db.query(Proveedor).count() == 0
    assert [f["estado"] for f in vista["filas"]] == ["lista", "lista", "lista", "omitida"]
    assert vista["fila_encabezado"] == 1
    assert vista["total_listas"] == "2478.99"
    assert vista["filas"][2]["datos"]["metodo_pago"] == "credito"
    creado = confirmar(client, auth_headers, vista)
    assert creado.status_code == 200, creado.text
    assert creado.json()["importados"] == 3
    assert creado.json()["monto_importado"] == "2478.99"
    assert db.query(Proveedor).count() == 2
    assert db.query(Egreso).filter_by(origen="excel").count() == 3
    eventos = db.query(LogAuditoria).filter_by(entidad="egreso", accion="crear").all()
    assert len(eventos) == 3
    audit = json.loads(eventos[0].datos_nuevos)["importacion_excel"]
    assert audit["archivo"] == "gastos.xlsx" and audit["hoja"] == "Septiembre" and audit["fila"] == 2
    assert len(audit["sha256"]) == 64
    retry = confirmar(client, auth_headers, vista)
    assert retry.status_code == 200, retry.text
    assert retry.json()["importados"] == 0 and retry.json()["ya_registrados"] == 3
    assert db.query(Egreso).count() == 3
    assert db.query(LogAuditoria).filter_by(entidad="egreso", accion="crear").count() == 3
    repetida = preview(client, auth_headers, contenido, nombre="renombrado.xlsx")
    assert [f["estado"] for f in repetida["filas"]][:3] == ["registrada"] * 3
    reporte = client.get("/api/v1/reportes/egresos?fecha_inicio=2026-09-01&fecha_fin=2026-09-30", headers=auth_headers)
    assert reporte.status_code == 200
    assert reporte.json()["resumen"]["total"] == 2478.99


def test_duplicados_manual_y_del_archivo_requieren_seleccion_explicita(client, auth_headers, db):
    db.add(Egreso(concepto="Gas LP", proveedor="Proveedor prueba", fecha=date(2026, 9, 22),
                  monto=1300, metodo_pago="efectivo", categoria="operativo"))
    db.commit()
    vista = preview(client, auth_headers, libro([
        ["22-sep-26", "Proveedor prueba", 1300, "EFECTIVO"],
        ["23-sep-26", "Molino prueba", 50, "TARJETA"],
        ["23-sep-26", "Molino prueba", 50, "TARJETA"],
    ]))
    assert [f["estado"] for f in vista["filas"]] == ["duplicado", "lista", "duplicado"]
    bloqueada = confirmar(client, auth_headers, vista, [2, 3, 4])
    assert bloqueada.status_code == 409
    assert db.query(Egreso).count() == 1
    permitida = confirmar(client, auth_headers, vista, [2, 3, 4], incluir_duplicados=[2, 4])
    assert permitida.status_code == 200, permitida.text
    assert permitida.json()["importados"] == 3


def test_duplicado_nuevo_entre_revision_y_confirmacion_revierte_todo(client, auth_headers, db):
    vista = preview(client, auth_headers, libro([
        ["01-sep-26", "Nuevo proveedor", 10, "EFECTIVO"],
        ["02-sep-26", "Ya registrado", 20, "EFECTIVO"],
    ]))
    db.add(Egreso(concepto="Manual", proveedor="Ya registrado", fecha=date(2026, 9, 2),
                  monto=20, metodo_pago="efectivo", categoria="operativo"))
    db.commit()
    response = confirmar(client, auth_headers, vista)
    assert response.status_code == 409, response.text
    assert db.query(Egreso).count() == 1
    assert db.query(Proveedor).count() == 0
    assert db.query(LogAuditoria).filter(LogAuditoria.entidad.in_(["egreso", "proveedor"])).count() == 0


def test_datos_invalidos_y_formulas_se_muestran_sin_inventar_valores(client, auth_headers, db):
    futura = (operation_today() + timedelta(days=31)).isoformat()
    vista = preview(client, auth_headers, libro([
        ["31-feb-26", "Proveedor", 10, "EFECTIVO"],
        [None, "Proveedor", 10, "EFECTIVO"],
        ["01-sep-26", "Proveedor", "=1+1", "EFECTIVO"],
        ["01-sep-26", "Proveedor", 10, "DESCONOCIDO"],
        ["01-sep-26", None, 10, "EFECTIVO"],
        [futura, "Proveedor", 10, "EFECTIVO"],
        ["01-sep-26", "Proveedor", 0, "EFECTIVO"],
        ["01-sep-26", "Proveedor", "USD 100", "EFECTIVO"],
        ["01-sep-26", "Proveedor", "NaN", "EFECTIVO"],
        ["01-sep-26", "Proveedor", 10, None],
    ]))
    assert all(f["estado"] == "error" and f["errores"] for f in vista["filas"])
    assert "Fecha invalida" in vista["filas"][0]["errores"][0]
    assert vista["token"] is None
    assert db.query(Egreso).count() == db.query(Proveedor).count() == 0


def test_mapeo_hojas_y_categoria_sin_proveedores_nuevos(client, auth_headers, db):
    contenido = libro([["2026-09-01", "Prov A", 40, "EFECTIVO"]], headers=["Dia", "Destino", "Valor", "Pago"],
        segunda=[["2026-09-03", "Prov B", 60, "TRANSFERENCIA"]])
    inicial = preview(client, auth_headers, contenido)
    assert inicial["faltantes"] == ["proveedor"]
    assert inicial["token"] is None
    mapeada = preview(client, auth_headers, contenido, {"fila_encabezado": 1,
        "columnas": {"fecha": 1, "proveedor": 2, "monto": 3, "metodo_pago": 4},
        "categoria": "insumos", "guardar_proveedor": False})
    assert mapeada["filas"][0]["datos"]["categoria"] == "insumos"
    assert confirmar(client, auth_headers, mapeada).status_code == 200
    assert db.query(Proveedor).count() == 0
    otra = preview(client, auth_headers, contenido, {"hoja": "Octubre"})
    assert otra["filas"][0]["datos"]["proveedor"] == "Prov B"
    assert otra["hojas"] == ["Septiembre", "Octubre"]


def test_sin_encabezados_con_total_y_fecha_serial_1904(client, auth_headers):
    wb = Workbook()
    wb.epoch = datetime(1904, 1, 1)
    from openpyxl.utils.datetime import to_excel
    wb.active.append([to_excel(datetime(2026, 9, 1), wb.epoch), "Proveedor", 15, "EFECTIVO"])
    wb.active.append(["Total", None, 15, None])
    contenido = BytesIO()
    wb.save(contenido)
    vista = preview(client, auth_headers, contenido.getvalue())
    assert vista["fila_encabezado"] == 0
    assert vista["filas"][0]["datos"]["fecha"] == "2026-09-01"
    assert vista["filas"][0]["estado"] == "lista"
    assert vista["filas"][1]["estado"] == "omitida"


def test_resguardo_lote_valida_acumulado_y_rollback(client, auth_headers, db, admin_user):
    db.add(ResguardoControl(id=1, saldo_inicial=100, ultimo_corte_id=0, ultimo_egreso_id=0,
        creado_por_id=admin_user.id, iniciado_en=datetime.now(timezone.utc) - timedelta(hours=1)))
    db.commit()
    fecha = operation_today().isoformat()
    vista = preview(client, auth_headers, libro([[fecha, "Proveedor A", 60, "RESGUARDO"],
                                               [fecha, "Proveedor B", 60, "RESGUARDO"]]))
    response = confirmar(client, auth_headers, vista)
    assert response.status_code == 400 and "insuficiente" in response.text
    assert db.query(Egreso).count() == db.query(Proveedor).count() == 0
    assert confirmar(client, auth_headers, vista, [2]).status_code == 200
    saldo = client.get("/api/v1/egresos/resguardo", headers=auth_headers).json()
    assert saldo["saldo"] == 40


def test_token_no_admite_manipulacion_cambio_usuario_o_caducidad(client, auth_headers, db, monkeypatch):
    vista = preview(client, auth_headers, libro([["01-sep-26", "Proveedor", 10, "EFECTIVO"]]))
    raw = vista["token"]
    assert confirmar(client, auth_headers, {**vista, "token": "x" + raw[1:]}).status_code == 400
    with pytest.raises(ValueError):
        svc.leer_vista_previa(raw, 999999)
    assert confirmar(client, auth_headers, vista, [999]).status_code == 400
    assert confirmar(client, auth_headers, vista, [2], incluir_duplicados=[2]).status_code == 400
    monkeypatch.setattr(svc.time, "time", lambda: 99999999999)
    assert confirmar(client, auth_headers, vista).status_code == 400
    assert db.query(Egreso).count() == 0


def test_importacion_rechaza_archivos_incorrectos_y_no_salta_limites(client, auth_headers, monkeypatch):
    for nombre, contenido in [("gastos.xls", b"not-xlsx"), ("foto.jpg", b"image"), ("invalido.xlsx", b"not-zip")]:
        response = client.post(PATH + "previsualizar", headers=auth_headers, files={"archivo": (nombre, contenido)})
        assert response.status_code == 400
    contenido = libro([["01-sep-26", "Proveedor", 10, "EFECTIVO"]] * 501)
    response = client.post(PATH + "previsualizar", headers=auth_headers, files={"archivo": ("grande.xlsx", contenido)})
    assert response.status_code == 400 and "500 gastos" in response.text
    monkeypatch.setattr(svc, "MAX_EXPANDED_BYTES", 10)
    response = client.post(PATH + "previsualizar", headers=auth_headers, files={"archivo": ("limitado.xlsx", libro([]))})
    assert response.status_code == 400


def test_permisos_igual_que_captura_manual(client, auth_headers, db, admin_user):
    contenido = libro([["01-sep-26", "Proveedor", 10, "EFECTIVO"]])
    assert client.post(PATH + "previsualizar", files={"archivo": ("gastos.xlsx", contenido)}).status_code == 401
    vista = preview(client, auth_headers, contenido)
    admin_user.rol = RolUsuario.CONSULTA
    db.commit()
    assert client.post(PATH + "previsualizar", headers=auth_headers, files={"archivo": ("gastos.xlsx", contenido)}).status_code == 403
    assert confirmar(client, auth_headers, vista).status_code == 403
    assert db.query(Egreso).count() == 0


@pytest.mark.parametrize("valor,esperado", [("$1,234.56", "1234.56"), ("1.234,56", "1234.56"),
    ("1,234", "1234.00"), ("23,50", "23.50"), (178.49000000000001, "178.49")])
def test_decimales_mxn_sin_flotantes_en_importes(valor, esperado):
    assert svc.monto_excel(valor) == esperado


@pytest.mark.parametrize("valor", ["1.23.4,56", "1,23,4.56", "10.001", "1.000", "Infinity", "-10", True])
def test_montos_mal_formados_no_se_corrigen_silenciosamente(valor):
    with pytest.raises(ValueError):
        svc.monto_excel(valor)


def test_500_gastos_mas_total_se_importan_sin_contar_dos_veces(client, auth_headers, db):
    contenido = libro([["01-sep-26", "Proveedor lote", i, "TARJETA"] for i in range(1, 501)] +
                      [["TOTAL", None, "=SUM(C2:C501)", None]])
    vista = preview(client, auth_headers, contenido)
    assert len(vista["filas"]) == 501
    assert vista["total_listas"] == "125250.00"
    response = confirmar(client, auth_headers, vista)
    assert response.status_code == 200, response.text
    assert response.json()["importados"] == 500
    assert db.query(Egreso).count() == 500 and db.query(Proveedor).count() == 1


def test_fechas_texto_americanas_son_opcion_explicita(client, auth_headers):
    contenido = libro([["09/01/2026", "Proveedor", 10, "EFECTIVO"]])
    mx = preview(client, auth_headers, contenido)
    us = preview(client, auth_headers, contenido, {"formato_fecha": "mes_dia_anio"})
    assert mx["filas"][0]["datos"]["fecha"] == "2026-01-09"
    assert us["filas"][0]["datos"]["fecha"] == "2026-09-01"


def test_importacion_anulada_no_se_reactiva_al_reintentar(client, auth_headers, db):
    contenido = libro([["01-sep-26", "Proveedor", 10, "EFECTIVO"]])
    vista = preview(client, auth_headers, contenido)
    assert confirmar(client, auth_headers, vista).status_code == 200
    egreso = db.query(Egreso).one()
    egreso.activo = False
    db.commit()
    nueva = preview(client, auth_headers, contenido)
    assert nueva["filas"][0]["estado"] == "registrada"
    retry = confirmar(client, auth_headers, vista)
    assert retry.status_code == 200 and retry.json()["anulados"] == 1
    assert db.query(Egreso).count() == 1 and db.query(Egreso).one().activo is False


def test_archivo_grande_y_opciones_invalidas_no_crean_gastos(client, auth_headers, db):
    grande = client.post(PATH + "previsualizar", headers=auth_headers,
        files={"archivo": ("grande.xlsx", b"0" * (svc.MAX_BYTES + 1))})
    assert grande.status_code == 413
    contenido = libro([["01-sep-26", "Proveedor", 10, "EFECTIVO"]])
    for opciones in [{"columnas": {"monto": 99}}, {"columnas": {"fecha": 1, "monto": 1}},
                     {"fila_encabezado": 50}, {"hoja": "No existe"}]:
        response = client.post(PATH + "previsualizar", headers=auth_headers,
            files={"archivo": ("gastos.xlsx", contenido)}, data={"opciones": json.dumps(opciones)})
        assert response.status_code == 400
    assert db.query(Egreso).count() == 0


def test_xml_con_entidades_se_rechaza_sin_ejecutar_formulas(client, auth_headers, db):
    original = libro([["01-sep-26", "Proveedor", 10, "EFECTIVO"]])
    salida = BytesIO()
    with ZipFile(BytesIO(original)) as src, ZipFile(salida, "w", compression=ZIP_DEFLATED) as dst:
        for info in src.infolist():
            contenido = src.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                contenido = b'<!DOCTYPE worksheet [<!ENTITY prueba "No se debe expandir">]>' + contenido
                contenido = contenido.replace(b">Proveedor<", b">&prueba;<")
            dst.writestr(info, contenido)
    response = client.post(PATH + "previsualizar", headers=auth_headers,
        files={"archivo": ("entidades.xlsx", salida.getvalue())})
    assert response.status_code == 400
    assert db.query(Egreso).count() == 0


def test_origen_excel_no_se_puede_fingir_en_la_captura_manual(client, auth_headers, db):
    response = client.post("/api/v1/egresos/", headers=auth_headers, json={
        "concepto": "Prueba", "proveedor": "Proveedor", "monto": "10.00", "origen": "excel"})
    assert response.status_code == 400
    assert db.query(Egreso).count() == 0

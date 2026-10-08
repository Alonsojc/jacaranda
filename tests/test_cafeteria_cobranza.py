"""Receivables stay independent of POS cash and retain the complete ledger."""

from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO
import re

import pytest
from openpyxl import load_workbook

from app.core.security import create_access_token, get_password_hash
from app.core.time_utils import operation_datetime, operation_now, operation_today
from app.models.cafeteria import (
    CafeteriaCliente, CafeteriaVenta, DetalleCafeteriaVenta,
    EstadoCuentaCafeteria, PagoCafeteriaVenta,
)
from app.models.inventario import Producto
from app.models.usuario import RolUsuario, Usuario
from app.models.venta import MetodoPago, TerminalPago
from app.services import cafeteria_export_service as exports


@pytest.fixture
def cartera(db, admin_user):
    a = CafeteriaCliente(nombre="Café Uno", activo=True)
    b = CafeteriaCliente(nombre="Café Dos", activo=False)
    p = Producto(codigo="COBRANZA-TEST", nombre="Panqué & chocolate <grande>",
                 precio_unitario=100, precio_cafeteria=80, stock_actual=20, stock_minimo=0)
    db.add_all([a, b, p])
    db.commit()
    return a, b, p, admin_user


def pedido(db, cartera, *, cliente=None, pagado=0, cancelado=False, legado=False, entrega=True):
    a, _, producto, user = cartera
    cafe = cliente or a
    estado = EstadoCuentaCafeteria.PAGADA if pagado >= 100 else (
        EstadoCuentaCafeteria.PARCIAL if pagado else EstadoCuentaCafeteria.PENDIENTE
    )
    venta = CafeteriaVenta(
        folio=f"CAF-TEST-{db.query(CafeteriaVenta).count() + 1:08d}",
        cafeteria_id=None if legado else cafe.id, cafeteria_nombre=cafe.nombre,
        usuario_id=user.id, subtotal=100, total=100, monto_pagado=pagado,
        fecha=operation_now() - timedelta(days=90),
        fecha_entrega=operation_today() - timedelta(days=10) if entrega else None,
        fecha_credito=operation_today() - timedelta(days=3), dias_credito=7,
        estado=EstadoCuentaCafeteria.CANCELADA if cancelado else estado,
        detalles=[DetalleCafeteriaVenta(producto_id=producto.id, cantidad=1,
                                       precio_unitario=100, subtotal=100, tasa_iva=0, monto_iva=0)],
    )
    db.add(venta)
    db.flush()
    if pagado:
        db.add(PagoCafeteriaVenta(venta_id=venta.id, monto=pagado,
               metodo_pago=MetodoPago.TRANSFERENCIA, terminal=TerminalPago.EFECTIVO,
               usuario_id=user.id, fecha=operation_now() - timedelta(days=2)))
    db.commit()
    return venta


def test_pendientes_filtra_cliente_inactivo_y_legado_sin_limite_de_semana(client, auth_headers, db, cartera):
    a, b, _, _ = cartera
    pendiente = pedido(db, cartera)
    pedido(db, cartera, pagado=100)
    pedido(db, cartera, cancelado=True)
    pedido(db, cartera, cliente=b)
    legado = pedido(db, cartera, legado=True, entrega=False)
    respuesta = client.get(f"/api/v1/cafeteria/ventas?pendientes=true&cafeteria_id={a.id}", headers=auth_headers)
    assert respuesta.status_code == 200, respuesta.text
    assert [v["id"] for v in respuesta.json()] == [pendiente.id]
    assert len(client.get("/api/v1/cafeteria/ventas?pendientes=true", headers=auth_headers).json()) == 3
    nombres = client.get("/api/v1/cafeteria/cobranza/clientes", headers=auth_headers).json()
    assert {"id": b.id, "nombre": b.nombre} in nombres
    assert {"id": None, "nombre": a.nombre} in nombres
    antiguos = client.get("/api/v1/cafeteria/ventas", params={"cafeteria_nombre": a.nombre}, headers=auth_headers).json()
    assert [v["id"] for v in antiguos] == [legado.id]
    assert antiguos[0]["fecha_entrega"] is None
    resumen = client.get(f"/api/v1/cafeteria/cobranza/resumen?cafeteria_id={a.id}", headers=auth_headers).json()
    assert resumen == {"cuentas": 1, "total": 100, "pagado": 0, "saldo": 100, "vencidas": 1}
    error = client.get(f"/api/v1/cafeteria/ventas?cafeteria_id={a.id}&cafeteria_nombre=otro", headers=auth_headers)
    assert error.status_code == 400


def test_abonos_fechados_reintento_y_liquidacion_conservan_inventario(client, auth_headers, db, cartera):
    venta = pedido(db, cartera)
    producto = cartera[2]
    stock = producto.stock_actual
    fecha = (operation_today() - timedelta(days=1)).isoformat()
    payload = {"idempotency_key": "abono-prueba-unico", "monto": "40.25", "fecha_pago": fecha,
               "metodo_pago": "04", "terminal": "clip", "referencia": "REF-01"}
    ruta = f"/api/v1/cafeteria/ventas/{venta.id}/pagos"
    for _ in range(2):
        resp = client.post(ruta, json=payload, headers=auth_headers)
        assert resp.status_code == 200, resp.text
        assert Decimal(resp.json()["monto_pagado"]) == Decimal("40.25")
        assert resp.json()["pagos"][0]["fecha_pago"] == fecha
    assert db.query(PagoCafeteriaVenta).filter_by(venta_id=venta.id).count() == 1
    conflicto = client.post(ruta, json={**payload, "monto": "30.00"}, headers=auth_headers)
    assert conflicto.status_code == 400
    otro = pedido(db, cartera)
    assert client.post(f"/api/v1/cafeteria/ventas/{otro.id}/pagos", json=payload, headers=auth_headers).status_code == 400
    liquidado = client.post(ruta, json={"monto": "59.75", "fecha_pago": fecha,
                            "idempotency_key": "liquidar-prueba-unico"}, headers=auth_headers)
    assert liquidado.status_code == 200, liquidado.text
    assert liquidado.json()["estado"] == "pagada"
    assert Decimal(liquidado.json()["saldo_pendiente"]) == 0
    assert len(liquidado.json()["pagos"]) == 2
    pendientes = client.get("/api/v1/cafeteria/ventas?pendientes=true", headers=auth_headers).json()
    assert venta.id not in [v["id"] for v in pendientes]
    historial = client.get("/api/v1/cafeteria/ventas", headers=auth_headers).json()
    assert venta.id in [v["id"] for v in historial]
    db.refresh(producto)
    assert producto.stock_actual == stock


@pytest.mark.parametrize("payload,status", [
    ({"monto": "100.01"}, 400), ({"monto": "0"}, 422),
    ({"monto": "1.001"}, 422), ({"monto": "NaN"}, 422),
    ({"monto": "1", "fecha_pago": "no-es-fecha"}, 422),
])
def test_pago_invalido_no_cambia_saldo(client, auth_headers, db, cartera, payload, status):
    venta = pedido(db, cartera)
    resp = client.post(f"/api/v1/cafeteria/ventas/{venta.id}/pagos", json=payload, headers=auth_headers)
    assert resp.status_code == status, resp.text
    db.refresh(venta)
    assert venta.monto_pagado == 0
    assert db.query(PagoCafeteriaVenta).count() == 0


def test_fechas_futuras_se_rechazan_y_entrega_corrige_credito_sin_caja(client, auth_headers, db, cartera):
    venta = pedido(db, cartera, entrega=False)
    futuro = (operation_today() + timedelta(days=1)).isoformat()
    for ruta, payload in [
        (f"/api/v1/cafeteria/ventas/{venta.id}/pagos", {"monto": "20", "fecha_pago": futuro}),
        (f"/api/v1/cafeteria/ventas/{venta.id}/entrega", {"fecha_entrega": futuro}),
    ]:
        response = client.request("POST" if "pagos" in ruta else "PUT", ruta, json=payload, headers=auth_headers)
        assert response.status_code == 400, response.text
    fecha = operation_today() - timedelta(days=5)
    resp = client.put(f"/api/v1/cafeteria/ventas/{venta.id}/entrega", json={"fecha_entrega": fecha.isoformat()}, headers=auth_headers)
    assert resp.status_code == 200, resp.text
    assert resp.json()["fecha_entrega"] == fecha.isoformat()
    assert resp.json()["fecha_credito"] == (fecha + timedelta(days=7)).isoformat()
    assert Decimal(resp.json()["total"]) == 100
    assert Decimal(resp.json()["monto_pagado"]) == 0
    assert cartera[2].stock_actual == 20


def test_nueva_entrega_fecha_y_abono_inicial_por_separado(client, auth_headers, db, cartera):
    a, _, p, _ = cartera
    entregado = operation_today() - timedelta(days=4)
    pagado = operation_today() - timedelta(days=2)
    resp = client.post("/api/v1/cafeteria/ventas", headers=auth_headers, json={
        "cafeteria_id": a.id, "cafeteria_nombre": a.nombre,
        "fecha_entrega": entregado.isoformat(), "fecha_pago_inicial": pagado.isoformat(),
        "pago_inicial": "25", "detalles": [{"producto_id": p.id, "cantidad": 1}],
    })
    assert resp.status_code == 201, resp.text
    venta = resp.json()
    assert venta["fecha_entrega"] == entregado.isoformat()
    assert venta["fecha_credito"] == (entregado + timedelta(days=7)).isoformat()
    assert venta["pagos"][0]["fecha_pago"] == pagado.isoformat()
    assert operation_datetime(db.get(CafeteriaVenta, venta["id"]).fecha).date() == operation_today()


def test_pdf_solo_deuda_cliente_escapa_nombres_y_no_escribe(client, auth_headers, db, cartera):
    a, b, _, _ = cartera
    pendiente = pedido(db, cartera, pagado=25)
    pedido(db, cartera, pagado=100)
    pedido(db, cartera, cancelado=True)
    pedido(db, cartera, cliente=b)
    antes = db.query(CafeteriaVenta).count(), db.query(PagoCafeteriaVenta).count(), cartera[2].stock_actual
    data = exports.estado_cuenta(db, a.id)
    assert data["cliente"] == a.nombre
    assert [v.id for v in data["ventas"]] == [pendiente.id]
    assert (data["total"], data["pagado"], data["saldo"]) == (100, 25, 75)
    resp = client.get(f"/api/v1/cafeteria/estado-cuenta/pdf?cafeteria_id={a.id}", headers=auth_headers)
    assert resp.status_code == 200, resp.text
    assert resp.content.startswith(b"%PDF")
    assert resp.headers["content-type"] == "application/pdf"
    assert "attachment" in resp.headers["content-disposition"]
    assert client.get("/api/v1/cafeteria/estado-cuenta/pdf", headers=auth_headers).status_code == 400
    assert client.get("/api/v1/cafeteria/estado-cuenta/pdf?cafeteria_id=999999", headers=auth_headers).status_code == 404
    assert antes == (db.query(CafeteriaVenta).count(), db.query(PagoCafeteriaVenta).count(), cartera[2].stock_actual)


def test_excel_completo_pagadas_canceladas_fechas_y_texto_no_formulas(client, auth_headers, db, cartera):
    pendiente = pedido(db, cartera, pagado=25)
    pagada = pedido(db, cartera, pagado=100)
    cancelada = pedido(db, cartera, pagado=10, cancelado=True, entrega=False)
    pendiente.notas = '=HYPERLINK("https://example.test", "nota")'
    db.commit()
    resp = client.get("/api/v1/cafeteria/historial/excel", headers=auth_headers)
    assert resp.status_code == 200, resp.text
    wb = load_workbook(BytesIO(resp.content))
    assert wb.sheetnames == ["Pedidos", "Partidas", "Pagos"]
    rows = list(wb["Pedidos"].values)
    assert len(rows) == 4
    assert {r[1] for r in rows[1:]} == {pendiente.folio, pagada.folio, cancelada.folio}
    assert {r[10] for r in rows[1:]} == {"parcial", "pagada", "cancelada"}
    assert rows[1][13:16] == (100, 25, 75)
    assert rows[3][15] == 0
    assert rows[3][7] is None
    assert wb["Pedidos"]["R2"].data_type == "s"
    assert wb["Pedidos"]["R2"].value.startswith("=HYPERLINK")
    assert isinstance(rows[1][6], datetime)
    assert wb["Pedidos"]["H2"].number_format == "yyyy-mm-dd"
    assert wb["Pagos"].max_row == 4
    assert sum(r[6] for r in list(wb["Pagos"].values)[1:]) == 135
    assert all(ws.freeze_panes == "A2" and ws.auto_filter.ref for ws in wb)


def test_historial_y_resumen_superan_500_registros_y_pdf_paginas(db, cartera):
    a, _, p, user = cartera
    ventas = [CafeteriaVenta(
        folio=f"CAF-MAS-{i:08d}", cafeteria_id=a.id, cafeteria_nombre=a.nombre,
        usuario_id=user.id, subtotal=100, total=100, monto_pagado=0,
        estado=EstadoCuentaCafeteria.PENDIENTE,
        detalles=[DetalleCafeteriaVenta(producto_id=p.id, cantidad=1, precio_unitario=100, subtotal=100)],
    ) for i in range(505)]
    db.add_all(ventas)
    db.commit()
    from app.services.cafeteria_service import listar_ventas, resumen_cobranza
    assert len(listar_ventas(db, limit=51, offset=500)) == 5
    assert resumen_cobranza(db, a.id)["saldo"] == 50500
    wb = load_workbook(exports.exportar_historial_excel(db))
    assert wb["Pedidos"].max_row == 506
    assert wb["Partidas"].max_row == 506
    pdf = exports.generar_estado_cuenta_pdf(db, a.id).getvalue()
    assert len(re.findall(rb"/Type /Page\b", pdf)) > 1


def test_pdf_pedido_con_muchas_partidas_puede_continuar_en_otra_pagina(db, cartera):
    venta = pedido(db, cartera)
    for _ in range(100):
        db.add(DetalleCafeteriaVenta(venta_id=venta.id, producto_id=cartera[2].id,
               cantidad=1, precio_unitario=100, subtotal=100))
    db.commit()
    pdf = exports.generar_estado_cuenta_pdf(db, cartera[0].id).getvalue()
    assert len(re.findall(rb"/Type /Page\b", pdf)) > 1


@pytest.mark.parametrize("path", ["/cobranza/clientes", "/cobranza/resumen", "/estado-cuenta/pdf", "/historial/excel"])
def test_descargas_y_cobranza_requieren_permiso(client, db, path):
    ruta = "/api/v1/cafeteria" + path
    assert client.get(ruta).status_code == 401
    user = Usuario(nombre="Sin acceso", email="sin-acceso@test.com",
                   hashed_password=get_password_hash("solo-test"), rol=RolUsuario.CAJERO,
                   permisos_modulos={"cafeteria": "oculto"})
    db.add(user)
    db.commit()
    token = create_access_token({"sub": str(user.id)})
    assert client.get(ruta, headers={"Authorization": "Bearer " + token}).status_code == 403


def test_permiso_ver_no_registra_pagos_ni_cambia_entrega(client, db, cartera):
    venta = pedido(db, cartera)
    user = Usuario(nombre="Consulta", email="consulta-cafe@test.com",
                   hashed_password=get_password_hash("solo-test"), rol=RolUsuario.CAJERO,
                   permisos_modulos={"cafeteria": "ver"})
    db.add(user)
    db.commit()
    headers = {"Authorization": "Bearer " + create_access_token({"sub": str(user.id)})}
    assert client.get("/api/v1/cafeteria/historial/excel", headers=headers).status_code == 200
    assert client.post(f"/api/v1/cafeteria/ventas/{venta.id}/pagos", json={"monto": 1}, headers=headers).status_code == 403
    assert client.put(f"/api/v1/cafeteria/ventas/{venta.id}/entrega", json={"fecha_entrega": operation_today().isoformat()}, headers=headers).status_code == 403
